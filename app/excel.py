from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import CABIN_LABELS, PROGRAM_LABELS

HEADERS = [
    "Origem",
    "Destino",
    "Ida",
    "Horário ida",
    "Chegada",
    "Volta",
    "Milhas",
    "Programa",
    "Taxas",
    "Companhia",
    "Operadoras",
    "Paradas",
    "Duração",
    "Conexões",
    "Tarifa",
    "Cabine",
    "Trecho",
    "Destaque",
    "Fonte",
    "Link",
    "Encontrado em",
]


def build_workbook(rows: list[dict]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Milhas"
    header_fill = PatternFill("solid", fgColor="1A2744")
    header_font = Font(color="F7F1E8", bold=True)
    for col, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, item in enumerate(rows, start=2):
        values = [
            item.get("origin"),
            item.get("destination"),
            item.get("departure_date"),
            item.get("departure_time") or "",
            item.get("arrival_time") or "",
            item.get("return_date") or "",
            item.get("miles"),
            PROGRAM_LABELS.get(item.get("miles_program") or "", item.get("miles_program") or ""),
            item.get("taxes"),
            item.get("airline") or "",
            item.get("operators") or "",
            item.get("stops"),
            item.get("duration") or "",
            item.get("layover") or "",
            item.get("fare") or "",
            CABIN_LABELS.get(item.get("cabin") or "", item.get("cabin") or ""),
            {"ida": "Só ida", "volta": "Só volta", "round_trip": "Ida + volta"}.get(
                item.get("trip_kind") or ("round_trip" if item.get("return_date") else "ida"),
                item.get("trip_kind") or "",
            ),
            "Melhor preço!" if item.get("best_price") else "",
            item.get("source") or "",
            item.get("booking_url") or "",
            item.get("found_at") or "",
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row_idx, col, value)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, len(rows) + 1)}"
    sheet.freeze_panes = "A2"
    widths = [10, 10, 12, 12, 12, 12, 12, 16, 10, 16, 10, 16, 14, 40, 22]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
